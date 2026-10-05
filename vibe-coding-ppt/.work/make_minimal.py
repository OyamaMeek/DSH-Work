from pptx import Presentation
from pptx.util import Inches, Pt
p = Presentation()
p.slide_width = Inches(13.333); p.slide_height = Inches(7.5)
s = p.slides.add_slide(p.slide_layouts[6])
tb = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.8))
r = tb.text_frame.paragraphs[0].add_run(); r.text = "Hello Render Test"; r.font.size = Pt(40)
p.save("vibe-coding-ppt/.work/minimal.pptx")
print("ok")
