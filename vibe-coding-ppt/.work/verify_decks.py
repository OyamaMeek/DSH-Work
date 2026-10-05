import sys
from pptx import Presentation
for path in sys.argv[1:]:
    prs = Presentation(path)
    print("==", path)
    print("   slides:", len(prs.slides), "size:", round(prs.slide_width/914400, 3), "x", round(prs.slide_height/914400, 3))
    nnotes = 0
    for i, s in enumerate(prs.slides, 1):
        has = bool(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
        nnotes += has
        texts = [sh.text_frame.text.replace("\n", " ") for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        print("   %02d notes=%s %s" % (i, "Y" if has else "-", (texts[0][:34] if texts else "")))
    print("   slides with notes:", nnotes)
