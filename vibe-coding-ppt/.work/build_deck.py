# -*- coding: utf-8 -*-
import json, sys, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

FONT = "PingFang SC"
BG      = RGBColor(0x0E, 0x17, 0x26)
CARD    = RGBColor(0x18, 0x28, 0x40)
CARD2   = RGBColor(0x14, 0x22, 0x38)
ACCENT  = RGBColor(0x22, 0xD3, 0xEE)
ACCENT2 = RGBColor(0xA7, 0x8B, 0xFA)
TEXT    = RGBColor(0xE8, 0xF0, 0xFA)
MUTED   = RGBColor(0x9F, 0xB3, 0xC8)
LINE    = RGBColor(0x2C, 0x40, 0x5C)
DARKTXT = RGBColor(0x08, 0x14, 0x22)

W = 13.333
H = 7.5


def set_font(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def style_run(run, size, color, bold=False, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_font(run, name)


def add_rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.0, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    return sp


def add_text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.15, space_after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    if space_after:
        p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    style_run(run, size, color, bold)
    return tb


def bg(slide):
    add_rect(slide, 0, 0, W, H, fill=BG)


def footer(slide, idx, total):
    add_text(slide, 0.75, H - 0.62, 6, 0.35, "Vibe Coding 学习系列", 10.5, MUTED)
    add_text(slide, W - 2.1, H - 0.62, 1.35, 0.35, "%02d / %02d" % (idx, total),
             10.5, MUTED, align=PP_ALIGN.RIGHT)


def slide_title(slide, text):
    add_text(slide, 0.75, 0.46, 11.8, 0.95, text, 30, TEXT, bold=True)
    add_rect(slide, 0.78, 1.38, 1.25, 0.065, fill=ACCENT)
    return 1.78


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def cover(slide, data, idx, total):
    add_rect(slide, 0, 0, 0.32, H, fill=ACCENT)
    add_rect(slide, 0.32, 0, 0.08, H, fill=ACCENT2)
    ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.2), Inches(4.1), Inches(5.4), Inches(5.4))
    ring.fill.background()
    ring.line.color.rgb = LINE
    ring.line.width = Pt(1.2)
    ring.shadow.inherit = False
    add_text(slide, 1.15, 2.05, 9.5, 0.4, data.get("kicker", ""), 15, ACCENT, bold=True)
    add_text(slide, 1.15, 2.65, 10.4, 1.7, data.get("title", ""), 46, TEXT, bold=True, line_spacing=1.08)
    add_rect(slide, 1.18, 4.55, 2.2, 0.08, fill=ACCENT)
    add_text(slide, 1.15, 4.95, 10.4, 1.2, data.get("subtitle", ""), 20, MUTED, line_spacing=1.25)
    add_text(slide, 1.15, H - 0.9, 8, 0.4, "根据 AI-Coding-Guide-Zh 教程整理", 12, MUTED)


def section(slide, data, idx, total):
    add_rect(slide, 0, 0, 0.18, H, fill=ACCENT)
    add_text(slide, 1.0, 2.6, 10.8, 1.2, data.get("title", ""), 40, TEXT, bold=True)
    add_text(slide, 1.03, 3.95, 10.8, 0.9, data.get("subtitle", ""), 19, MUTED, line_spacing=1.25)
    add_rect(slide, 1.03, 2.35, 1.6, 0.07, fill=ACCENT2)
    footer(slide, idx, total)


def bullets(slide, data, idx, total):
    top = slide_title(slide, data.get("title", ""))
    items = data.get("bullets", [])
    n = max(len(items), 1)
    avail = 7.15 - top - 0.35
    gap = min(avail / n, 1.22)
    size = 19 if n <= 4 else (17 if n == 5 else 15.5)
    longest = max((len(str(i)) for i in items), default=0)
    if longest > 46:
        size = min(size, 15.5)
    for i, item in enumerate(items):
        y = top + 0.12 + i * gap
        add_rect(slide, 0.82, y + 0.13, 0.17, 0.17, fill=ACCENT, shape=MSO_SHAPE.OVAL)
        add_text(slide, 1.22, y, 11.3, gap - 0.02, str(item), size, TEXT, line_spacing=1.22)
    footer(slide, idx, total)
    notes(slide, data.get("notes"))


def cards(slide, data, idx, total):
    top = slide_title(slide, data.get("title", ""))
    items = data.get("cards", [])
    n = len(items)
    cols = 2 if n >= 4 else (3 if n == 3 else 2)
    rows = (n + cols - 1) // cols
    gx, gy = 0.35, 0.32
    total_w = 11.85
    cw = (total_w - gx * (cols - 1)) / cols
    ch = min((7.05 - top - gy * (rows - 1)) / rows, 2.45)
    for i, c in enumerate(items):
        r, col = divmod(i, cols)
        x = 0.75 + col * (cw + gx)
        y = top + r * (ch + gy)
        add_rect(slide, x, y, cw, ch, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
        add_rect(slide, x, y, 0.075, ch, fill=ACCENT)
        add_text(slide, x + 0.34, y + 0.28, cw - 0.6, 0.5, str(c.get("title", "")), 18, ACCENT, bold=True)
        add_text(slide, x + 0.34, y + 0.88, cw - 0.6, ch - 1.1, str(c.get("desc", "")), 14, TEXT, line_spacing=1.28)
    footer(slide, idx, total)
    notes(slide, data.get("notes"))


def table(slide, data, idx, total):
    top = slide_title(slide, data.get("title", ""))
    headers = data.get("headers", [])
    rows = data.get("rows", [])
    nrow = len(rows) + 1
    ncol = len(headers)
    col_w = data.get("colWidths")
    total_w = 11.85
    height = min(0.46 * nrow + 0.1, 7.05 - top)
    shape = slide.shapes.add_table(nrow, ncol, Inches(0.75), Inches(top + 0.05), Inches(total_w), Inches(height))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    if not col_w:
        weights = []
        for c in range(ncol):
            cells = [str(headers[c])] + [str(r[c]) if c < len(r) else "" for r in rows]
            weights.append(max(len(s) for s in cells))
        weights = [max(w, 4) ** 0.75 for w in weights]
        ssum = sum(weights)
        col_w = [max(w / ssum, 0.13) for w in weights]
        ssum = sum(col_w)
        col_w = [f / ssum for f in col_w]
    for i, frac in enumerate(col_w):
        tbl.columns[i].width = Emu(int(Inches(total_w) * frac))
    tbl.rows[0].height = Inches(0.5)
    for r in range(1, nrow):
        tbl.rows[r].height = Inches(0.44)
    for c, htxt in enumerate(headers):
        set_cell(tbl.cell(0, c), htxt, 14.5, DARKTXT, bold=True, fill=ACCENT)
    for r, row in enumerate(rows, start=1):
        fill = CARD if r % 2 == 1 else CARD2
        for c in range(ncol):
            val = row[c] if c < len(row) else ""
            set_cell(tbl.cell(r, c), val, 13.5, TEXT, fill=fill)
    footer(slide, idx, total)
    notes(slide, data.get("notes"))


def set_cell(cell, text, size, color, bold=False, fill=None):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(0.12)
    cell.margin_right = Inches(0.1)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(text)
    for run in p.runs:
        style_run(run, size, color, bold)


def steps(slide, data, idx, total):
    top = slide_title(slide, data.get("title", ""))
    items = data.get("steps", [])
    n = max(len(items), 1)
    gy = 0.35
    w = (11.85 - gy * (n - 1)) / n
    y = top + 0.55
    h = 3.6
    for i, s in enumerate(items):
        x = 0.75 + i * (w + gy)
        add_rect(slide, x, y, w, h, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        add_rect(slide, x, y, w, 0.075, fill=ACCENT2)
        add_text(slide, x + 0.26, y + 0.45, w - 0.5, 0.6, str(i + 1), 34, ACCENT, bold=True)
        add_text(slide, x + 0.26, y + 1.2, w - 0.5, 0.7, str(s.get("title", "")), 17.5, TEXT, bold=True, line_spacing=1.15)
        add_text(slide, x + 0.26, y + 1.95, w - 0.5, h - 2.2, str(s.get("desc", "")), 13.5, MUTED, line_spacing=1.3)
    footer(slide, idx, total)
    notes(slide, data.get("notes"))


def quote(slide, data, idx, total):
    add_rect(slide, 0, 0, 0.18, H, fill=ACCENT2)
    add_text(slide, 1.1, 2.2, 2.0, 1.2, "\u201c", 90, ACCENT, bold=True)
    add_text(slide, 1.15, 3.15, 10.9, 1.9, data.get("title", ""), 28, TEXT, bold=True, line_spacing=1.35)
    add_text(slide, 1.18, 5.25, 10.8, 0.6, data.get("subtitle", ""), 16, MUTED)
    footer(slide, idx, total)


def end(slide, data, idx, total):
    add_rect(slide, 0, 0, W, H, fill=BG)
    add_rect(slide, W / 2 - 0.9, 2.55, 1.8, 0.08, fill=ACCENT)
    add_text(slide, 1.5, 2.95, 10.33, 1.2, data.get("title", ""), 38, TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.5, 4.35, 10.33, 0.9, data.get("subtitle", ""), 18, MUTED, align=PP_ALIGN.CENTER, line_spacing=1.3)
    add_text(slide, 1.5, H - 1.15, 10.33, 0.5, "AI-Coding-Guide-Zh · 老金", 12.5, MUTED, align=PP_ALIGN.CENTER)


RENDER = {"cover": cover, "section": section, "bullets": bullets, "cards": cards,
          "table": table, "steps": steps, "quote": quote, "end": end}


def build(json_path, out_path):
    with open(json_path, encoding="utf-8") as f:
        data = json.load(f)
    slides = data["slides"]
    total = len(slides)
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    for i, sd in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        bg(slide)
        layout = sd.get("layout", "bullets")
        fn = RENDER.get(layout)
        if fn is None:
            raise SystemExit("unknown layout: %s in %s" % (layout, json_path))
        if layout in ("cover", "section", "quote", "end"):
            fn(slide, sd, i, total)
        else:
            fn(slide, sd, i, total)
    prs.save(out_path)
    print("built %s -> %s (%d slides)" % (json_path, out_path, total))


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
